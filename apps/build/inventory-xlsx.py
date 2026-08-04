#!/usr/bin/env python3
"""
inventory-xlsx.py — z inventory-data.mjs vygeneruje apps/prehlad-aplikacii.xlsx

5 listov: Prehľad · Funkcie · Obrazovky · Náklady · Riziká.
Súčty a počty sú formuly (COUNTIF / SUMIF / SUM), nie dopočítané v Pythone —
po zmene dát sa hárok prepočíta sám.

Použitie:
    node inventory-data.mjs > /tmp/data.json
    python3 inventory-xlsx.py /tmp/data.json ../prehlad-aplikacii.xlsx
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ── Aura paleta (tlačová/svetlá varianta, aby bol hárok čitateľný aj na papieri) ──
TEAL = "03797E"
TEAL_L = "E3F4F4"
GOLD = "8A6417"
INK = "131B1A"
INK_3 = "5A6968"
LINE = "DED5DC"
BG_HD = "0E1413"
OK_BG, OK_FG = "E4F6EC", "1E7A46"
CO_BG, CO_FG = "FDF1DC", "8A5A0B"
NO_BG, NO_FG = "FBE7E5", "A32A22"
Q_BG, Q_FG = "F1EEF1", "5A6968"

FONT = "Arial"
H1 = Font(name=FONT, size=15, bold=True, color=INK)
H2 = Font(name=FONT, size=11, bold=True, color=TEAL)
TH = Font(name=FONT, size=9.5, bold=True, color="FFFFFF")
BODY = Font(name=FONT, size=10, color=INK)
BODY_B = Font(name=FONT, size=10, bold=True, color=INK)
SMALL = Font(name=FONT, size=9, color=INK_3)
MONO_SM = Font(name="Consolas", size=9, color=INK_3)
TOTAL = Font(name=FONT, size=10, bold=True, color=INK)

FILL_TH = PatternFill("solid", fgColor=BG_HD)
FILL_TOTAL = PatternFill("solid", fgColor=TEAL_L)
FILL_ZEBRA = PatternFill("solid", fgColor="FAF8FA")

THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_ = Border(top=Side(style="medium", color=TEAL))

WRAP_T = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
CTR = Alignment(horizontal="center", vertical="top")
RIGHT = Alignment(horizontal="right", vertical="top")

EUR = '#,##0\\ "€"'

BADGE = {
    "overené": (OK_BG, OK_FG),
    "rekonštruované": (CO_BG, CO_FG),
    "ukážkové": (Q_BG, Q_FG),
    "aktívna": (OK_BG, OK_FG),
    "nevyužitá": (CO_BG, CO_FG),
    "vysoké": (NO_BG, NO_FG),
    "stredné": (CO_BG, CO_FG),
    "nízke": (Q_BG, Q_FG),
    "V prevádzke": (OK_BG, OK_FG),
    "MVP v prevádzke": (OK_BG, OK_FG),
    "Appka neexistuje": (NO_BG, NO_FG),
}


def title_block(ws, title, subtitle, width):
    """Nadpis listu + podtitulok, vráti prvý voľný riadok."""
    ws["A1"] = title
    ws["A1"].font = H1
    ws["A2"] = subtitle
    ws["A2"].font = SMALL
    ws["A2"].alignment = WRAP_T
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 26
    return 4


def header_row(ws, row, headers):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = TH
        cell.fill = FILL_TH
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def write_row(ws, row, values, *, widths=None, badges=(), bold_first=True, money=(), mono=()):
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.border = BOX
        cell.alignment = WRAP_T
        cell.font = BODY
        if c == 1 and bold_first:
            cell.font = BODY_B
        if c in mono:
            cell.font = MONO_SM
        if c in money:
            cell.number_format = EUR
            cell.alignment = RIGHT
        if c in badges and v in BADGE:
            bg, fg = BADGE[v]
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.font = Font(name=FONT, size=10, bold=True, color=fg)
            cell.alignment = CTR


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_table(ws, name, first_row, last_row, last_col):
    """Autofilter + pruhovanie ako Excel tabuľka."""
    ref = f"A{first_row}:{get_column_letter(last_col)}{last_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight8", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(t)


def main():
    data_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])
    D = json.loads(data_path.read_text(encoding="utf-8"))
    META, MODULY, SPOLOCNE = D["META"], D["MODULY"], D["SPOLOCNE"]
    NAKLADY, RIZIKA, GLOB = D["NAKLADY"], D["RIZIKA"], D["GLOBALNE_OBRAZOVKY"]

    wb = Workbook()

    # ═════════════════ 1. Funkcie (naplní sa prvý, Prehľad naň odkazuje) ═════════════════
    ws_fn = wb.create_sheet("Funkcie")
    r = title_block(
        ws_fn,
        "Funkcie po moduloch",
        "Jeden riadok = jedna funkcia. Stĺpce Zdroj a Overenie hovoria, odkiaľ údaj je a či sa "
        "dá v repe skontrolovať. Filtruj cez šípky v hlavičke.",
        7,
    )
    header_row(ws_fn, r, ["Modul", "Reálna appka", "Funkcia", "Kategória", "Platí pre", "Zdroj", "Overenie"])
    fn_first = r + 1
    row = fn_first
    for m in MODULY:
        for f in m["funkcie"]:
            write_row(
                ws_fn,
                row,
                [m["nazov"], m["realnaAppka"], f[0], f[1], m["nazov"], f[2], f[3]],
                badges={7},
                mono={6},
            )
            row += 1
    for s in SPOLOCNE:
        write_row(
            ws_fn,
            row,
            ["Spoločná vrstva", "mockup Aura Suite", f"{s[0]} — {s[1]}", "Spoločné", s[2], s[3], s[4]],
            badges={7},
            mono={6},
        )
        row += 1
    fn_last = row - 1
    set_widths(ws_fn, [17, 26, 72, 14, 24, 30, 16])
    add_table(ws_fn, "Funkcie", r, fn_last, 7)

    # ═════════════════ 2. Obrazovky ═════════════════
    ws_sc = wb.create_sheet("Obrazovky")
    r = title_block(
        ws_sc,
        "Obrazovky",
        "Obrazovka = samostatná hash-routa v mockupe apps/aura-apps-hub.html. Modulové obrazovky "
        "vrátane intro stránky modulu; vstupné obrazovky a rozcestník sú vedené ako Globálne.",
        6,
    )
    header_row(ws_sc, r, ["Modul", "Reálna appka", "#", "Routa", "Obrazovka", "Typ / obsah"])
    sc_first = r + 1
    row = sc_first
    for mi, m in enumerate(MODULY, start=1):
        for oi, o in enumerate(m["obrazovky"], start=1):
            write_row(
                ws_sc,
                row,
                [m["nazov"], m["realnaAppka"], f"{mi}.{oi}", f"#{m['id']}/{o[0]}", o[1], o[2]],
                mono={3, 4},
            )
            row += 1
    for gi, g in enumerate(GLOB, start=1):
        write_row(
            ws_sc,
            row,
            ["Globálne", "mockup Aura Suite", f"0.{gi}", f"#{g[0]}", g[1], g[2]],
            mono={3, 4},
        )
        row += 1
    sc_last = row - 1
    set_widths(ws_sc, [17, 26, 7, 22, 24, 46])
    add_table(ws_sc, "Obrazovky", r, sc_last, 6)

    row += 1
    ws_sc.cell(row=row, column=1, value="Spolu obrazoviek").font = TOTAL
    c = ws_sc.cell(row=row, column=3, value=f"=COUNTA(E{sc_first}:E{sc_last})")
    c.font = TOTAL
    c.fill = FILL_TOTAL
    c.border = BOX
    ws_sc.cell(
        row=row,
        column=5,
        value=f'=COUNTIF(A{sc_first}:A{sc_last},"Globálne")&" globálnych"',
    ).font = SMALL

    # ═════════════════ 3. Náklady ═════════════════
    ws_co = wb.create_sheet("Náklady")
    r = title_block(
        ws_co,
        "Náklady",
        "Vlastné Aura appky nemajú licenčný náklad — bežia na vlastnom hardware v Dockeri. "
        "Register externých SaaS je z obrazovky HR → Aplikácie; sumy sú tam označené ako čiastočne ukážkové.",
        7,
    )
    ws_co.cell(row=r, column=1, value="Externé SaaS aplikácie").font = H2
    r += 1
    header_row(ws_co, r, ["Aplikácia", "Účel", "Vlastník", "€ / mes.", "Licencie", "Stav", "Overenie"])
    saas_first = r + 1
    row = saas_first
    for s in NAKLADY["saas"]:
        write_row(
            ws_co,
            row,
            [s[0], s[1], s[2], s[3], s[4], s[5], "ukážkové"],
            badges={6, 7},
            money={4},
        )
        row += 1
    saas_last = row - 1

    # súčtový riadok — formuly
    ws_co.cell(row=row, column=1, value="Spolu / mes.").font = TOTAL
    for col, formula in (
        (4, f"=SUM(D{saas_first}:D{saas_last})"),
        (5, f"=SUM(E{saas_first}:E{saas_last})"),
    ):
        cell = ws_co.cell(row=row, column=col, value=formula)
        cell.font = TOTAL
        cell.fill = FILL_TOTAL
        cell.border = BOX
        cell.number_format = EUR if col == 4 else "#,##0"
        cell.alignment = RIGHT
    cell = ws_co.cell(row=row, column=3, value=f"=COUNTA(A{saas_first}:A{saas_last})&\" aplikácií\"")
    cell.font = TOTAL
    cell.fill = FILL_TOTAL
    cell.border = BOX
    total_row = row
    row += 1

    ws_co.cell(row=row, column=1, value="Z toho nevyužité").font = BODY_B
    cell = ws_co.cell(row=row, column=4, value=f'=SUMIF(F{saas_first}:F{saas_last},"nevyužitá",D{saas_first}:D{saas_last})')
    cell.font = BODY_B
    cell.number_format = EUR
    cell.alignment = RIGHT
    cell.border = BOX
    cell = ws_co.cell(row=row, column=6, value=f"=D{row}/D{total_row}")
    cell.font = BODY_B
    cell.number_format = "0.0%"
    cell.alignment = CTR
    cell.border = BOX
    ws_co.cell(row=row, column=2, value="potenciál úspory podľa appky (Adobe CC + Canva Pro)").font = SMALL
    row += 2

    ws_co.cell(row=row, column=1, value="Interné rozpočty a prevádzkové náklady").font = H2
    row += 1
    header_row(ws_co, row, ["Položka", "Popis", "Vlastník", "Suma", "Obdobie", "Poznámka", "Overenie"])
    row += 1
    for i in NAKLADY["interne"]:
        write_row(ws_co, row, [i[0], i[1], i[2], i[3], i[4], i[5], "ukážkové"], badges={7}, money={4})
        row += 1

    row += 1
    for p in NAKLADY["poznamky"]:
        ws_co.cell(row=row, column=1, value=f"• {p}").font = SMALL
        ws_co.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws_co.row_dimensions[row].height = 24
        ws_co.cell(row=row, column=1).alignment = WRAP_T
        row += 1
    set_widths(ws_co, [24, 52, 14, 12, 12, 30, 16])

    # ═════════════════ 4. Riziká ═════════════════
    ws_ri = wb.create_sheet("Riziká")
    r = title_block(
        ws_ri,
        "Riziká a otvorené body",
        "Vychádzajú z reconu reálnych aplikácií (Hades pamäť) — v tomto repe ich nevidno, preto sú "
        "vedené ako rekonštruované. Zoradené od najzávažnejších.",
        7,
    )
    header_row(ws_ri, r, ["Modul", "Riziko / otvorený bod", "Dopad", "Závažnosť", "Stav", "Kontext a zmiernenie", "Zdroj"])
    ri_first = r + 1
    row = ri_first
    for x in RIZIKA:
        write_row(
            ws_ri,
            row,
            [x[0], x[1], x[2], x[3], x[4], x[6], x[5]],
            badges={4},
            mono={7},
        )
        row += 1
    ri_last = row - 1
    set_widths(ws_ri, [14, 46, 52, 13, 18, 52, 26])
    add_table(ws_ri, "Rizika", r, ri_last, 7)

    row += 1
    ws_ri.cell(row=row, column=1, value="Spolu").font = TOTAL
    for col, formula, fmt in (
        (2, f"=COUNTA(B{ri_first}:B{ri_last})&\" rizík\"", None),
        (4, f'=COUNTIF(D{ri_first}:D{ri_last},"vysoké")&" vysokých"', None),
    ):
        cell = ws_ri.cell(row=row, column=col, value=formula)
        cell.font = TOTAL
        cell.fill = FILL_TOTAL
        cell.border = BOX

    # ═════════════════ 5. Prehľad (krycí list, prvý v poradí) ═════════════════
    ws = wb.active
    ws.title = "Prehľad"
    ws["A1"] = META["titul"]
    ws["A1"].font = Font(name=FONT, size=17, bold=True, color=INK)
    ws["A2"] = f'{META["podtitul"]} · stav k {META["datum"]}'
    ws["A2"].font = Font(name=FONT, size=10, color=TEAL)
    ws["A3"] = META["rozsah"]
    ws["A3"].font = SMALL
    ws.row_dimensions[1].height = 24

    row = 5
    ws.cell(row=row, column=1, value="Súhrn").font = H2
    row += 1
    # očakávané výsledky, dopočítané v Pythone z tých istých dát — slúžia ako
    # kontrola formúl a ako viditeľná hodnota v prehliadačoch, ktoré formuly nepočítajú
    fn_all = [f[3] for m in MODULY for f in m["funkcie"]] + [s[4] for s in SPOLOCNE]
    exp = {
        "modulov": len(MODULY),
        "funkcii": len(fn_all),
        "overene": fn_all.count("overené"),
        "rekon": fn_all.count("rekonštruované"),
        "ukazk": fn_all.count("ukážkové"),
        "obrazovky": sum(len(m["obrazovky"]) for m in MODULY) + len(GLOB),
        "rizika": len(RIZIKA),
        "vysoke": sum(1 for x in RIZIKA if x[3] == "vysoké"),
        "saas": sum(s[3] for s in NAKLADY["saas"]),
        "uspora": sum(s[3] for s in NAKLADY["saas"] if s[5] == "nevyužitá"),
    }

    header_row(ws, row, ["Ukazovateľ", "Hodnota", "Ako sa počíta", "Kontrola"])
    row += 1
    sumar = [
        ("Aplikácií / modulov", f"=COUNTA(Prehľad!A{{MOD_FIRST}}:A{{MOD_LAST}})", "riadky v tabuľke Moduly nižšie", exp["modulov"]),
        ("Funkcií spolu", f"=COUNTA(Funkcie!C{fn_first}:C{fn_last})", "list Funkcie", exp["funkcii"]),
        ("— z toho overených", f'=COUNTIF(Funkcie!G{fn_first}:G{fn_last},"overené")', 'Funkcie · Overenie = "overené"', exp["overene"]),
        ("— rekonštruovaných", f'=COUNTIF(Funkcie!G{fn_first}:G{fn_last},"rekonštruované")', 'Funkcie · Overenie = "rekonštruované"', exp["rekon"]),
        ("— ukážkových", f'=COUNTIF(Funkcie!G{fn_first}:G{fn_last},"ukážkové")', 'Funkcie · Overenie = "ukážkové"', exp["ukazk"]),
        ("Obrazoviek spolu", f"=COUNTA(Obrazovky!E{sc_first}:E{sc_last})", "list Obrazovky", exp["obrazovky"]),
        ("Rizík spolu", f"=COUNTA(Riziká!B{ri_first}:B{ri_last})", "list Riziká", exp["rizika"]),
        ("— s vysokou závažnosťou", f'=COUNTIF(Riziká!D{ri_first}:D{ri_last},"vysoké")', 'Riziká · Závažnosť = "vysoké"', exp["vysoke"]),
        ("Externé SaaS — € / mes.", f"=SUM(Náklady!D{saas_first}:D{saas_last})", "list Náklady", exp["saas"]),
        ("— potenciál úspory", f'=SUMIF(Náklady!F{saas_first}:F{saas_last},"nevyužitá",Náklady!D{saas_first}:D{saas_last})', "aplikácie označené ako nevyužité", exp["uspora"]),
    ]
    sum_first = row
    for label, formula, how, control in sumar:
        write_row(ws, row, [label, None, how, control], bold_first=True)
        c = ws.cell(row=row, column=4)
        c.font = SMALL
        c.alignment = CTR
        cell = ws.cell(row=row, column=2, value=formula)
        cell.font = TOTAL
        cell.border = BOX
        cell.alignment = RIGHT
        if "€" in label or "úspory" in label:
            cell.number_format = EUR
        row += 1
    sum_last = row - 1

    row += 1
    ws.cell(row=row, column=1, value="Moduly").font = H2
    row += 1
    mod_hd = [
        "Modul", "Účel", "Reálna appka", "Port / URL", "Stav", "Hub /api/summary",
        "Vetva", "Stack", "Testy", "Git", "Veľkosť", "Kontajnery", "Vlastník",
        "Náklady", "Funkcií", "Obrazoviek", "Rizík",
    ]
    header_row(ws, row, mod_hd)
    row += 1
    mod_first = row
    for m in MODULY:
        krat = m["nazov"].replace("Aura ", "")
        write_row(
            ws,
            row,
            [
                m["nazov"], m["ucel"], m["realnaAppka"], m["port"], m["stav"]["label"],
                m["hubApi"]["label"], m["vetva"], m["stack"], m["testy"], m["git"],
                m["velkost"], m["kontajnery"], m["vlastnik"], m["naklady"],
                None, None, None,
            ],
            badges={5},
            mono={4},
        )
        for col, formula in (
            (15, f'=COUNTIF(Funkcie!A{fn_first}:A{fn_last},A{row})'),
            (16, f'=COUNTIF(Obrazovky!A{sc_first}:A{sc_last},A{row})'),
            (17, f'=COUNTIF(Riziká!A{ri_first}:A{ri_last},"{krat}")'),
        ):
            cell = ws.cell(row=row, column=col, value=formula)
            cell.font = TOTAL
            cell.border = BOX
            cell.alignment = CTR
        row += 1
    mod_last = row - 1

    ws.cell(row=row, column=1, value="Spolu").font = TOTAL
    for col in (15, 16, 17):
        L = get_column_letter(col)
        cell = ws.cell(row=row, column=col, value=f"=SUM({L}{mod_first}:{L}{mod_last})")
        cell.font = TOTAL
        cell.fill = FILL_TOTAL
        cell.border = BOX
        cell.alignment = CTR
    cell = ws.cell(
        row=row,
        column=2,
        value=(
            f'Súčet je len za 6 modulov. Naviac: spoločná vrstva ({len(SPOLOCNE)} funkcií) v liste Funkcie, '
            f'globálne obrazovky ({len(GLOB)}) v liste Obrazovky a riziká vedené ako „Spoločné" v liste Riziká.'
        ),
    )
    cell.font = SMALL
    cell.alignment = WRAP_T
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=14)
    row += 2

    # doplň odkaz na tabuľku Moduly do súhrnu
    ws.cell(row=sum_first, column=2).value = f"=COUNTA(A{mod_first}:A{mod_last})"

    ws.cell(row=row, column=1, value="Legenda — stupeň overenia").font = H2
    row += 1
    for k, v in D["LEGENDA"]["overenie"]:
        write_row(ws, row, [k, v], badges={1}, bold_first=False)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Zdroje").font = H2
    row += 1
    for label, val in (
        ("repo", META["zdrojRepo"]),
        ("Hades", META["zdrojHades"]),
        ("Poznámka", META["poznamka"]),
        ("Prestavba", "node apps/build/inventory-data.mjs > data.json && python3 apps/build/inventory-xlsx.py data.json apps/prehlad-aplikacii.xlsx"),
    ):
        ws.cell(row=row, column=1, value=label).font = BODY_B
        cell = ws.cell(row=row, column=2, value=val)
        cell.font = SMALL
        cell.alignment = WRAP_T
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.row_dimensions[row].height = 26
        row += 1

    set_widths(ws, [26, 44, 26, 20, 17, 20, 20, 60, 34, 40, 40, 30, 20, 34, 10, 12, 9])
    ws.freeze_panes = "A6"

    # poradie listov: Prehľad, Funkcie, Obrazovky, Náklady, Riziká
    wb._sheets = [ws, ws_fn, ws_sc, ws_co, ws_ri]
    for s in wb:
        s.sheet_view.showGridLines = False
        s.sheet_properties.tabColor = TEAL

    # openpyxl zapisuje formuly bez uložených výsledkov — toto prinúti Excel
    # (aj LibreOffice a Google Sheets) prepočítať celý hárok pri otvorení.
    wb.calculation.fullCalcOnLoad = True

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"✓ {out_path}")

    # ── kontrola: čo majú formuly vypočítať (dopočítané v Pythone z tých istých dát) ──
    fn_all = [f for m in MODULY for f in m["funkcie"]] + [[s[0], "Spoločné", s[3], s[4]] for s in SPOLOCNE]
    over = {k: sum(1 for f in fn_all if f[3] == k) for k in ("overené", "rekonštruované", "ukážkové")}
    saas_total = sum(s[3] for s in NAKLADY["saas"])
    saas_unused = sum(s[3] for s in NAKLADY["saas"] if s[5] == "nevyužitá")
    screens = sum(len(m["obrazovky"]) for m in MODULY) + len(GLOB)
    print(
        "  očakávané výsledky formúl: "
        f"modulov {len(MODULY)} · funkcií {len(fn_all)} "
        f"(overených {over['overené']}, rekonštruovaných {over['rekonštruované']}, ukážkových {over['ukážkové']}) · "
        f"obrazoviek {screens} · rizík {len(RIZIKA)} "
        f"(vysokých {sum(1 for x in RIZIKA if x[3] == 'vysoké')}) · "
        f"SaaS {saas_total} € (nevyužité {saas_unused} €)"
    )
    per_modul = {m["nazov"]: (len(m["funkcie"]), len(m["obrazovky"])) for m in MODULY}
    print("  na modul (funkcií / obrazoviek): " + " · ".join(f"{k} {v[0]}/{v[1]}" for k, v in per_modul.items()))


if __name__ == "__main__":
    main()
